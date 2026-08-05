"""Inspecciona cabeceras reales: celdas filas 1-12 + posición de imágenes."""
from __future__ import annotations

import zipfile
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

FILES = {
    "detector_metales": r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\comprobación detector de metales.xlsm",
    "pales": r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Supervisión del estado de los palés R2.xlsm",
    "lava_utiles": r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Registro_LavaUtiles.xlsm",
    "marcas_v2": r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Registro de marcas propias V2 .xlsm",
    "transporte": r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Supervisión de la calidad del transporte R4.xlsm",
    "cuchillos": r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Control de cuchillos produccion R3.xlsm",
    "cutters": r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Control de cutters R4.xlsm",
}

OUT = Path(r"C:\Users\ACZ\.codex\REGISTROS_DE_SALA\plantillas\_analisis\headers_detail.txt")
lines = []

for key, path in FILES.items():
    lines.append("=" * 72)
    lines.append(key)
    wb = load_workbook(path, keep_vba=True, data_only=False)
    ws = wb.worksheets[0]
    lines.append(f"sheet={ws.title!r} merges={len(ws.merged_cells.ranges)}")
    # first merges
    for m in list(ws.merged_cells.ranges)[:25]:
        lines.append(f"  merge {m}")
    # cells rows 1-12
    for r in range(1, 13):
        row_vals = []
        for c in range(1, min(12, (ws.max_column or 1)) + 1):
            cell = ws.cell(r, c)
            if cell.value is not None:
                fill = None
                if cell.fill and cell.fill.patternType:
                    fill = getattr(cell.fill.fgColor, "rgb", None)
                row_vals.append(f"{get_column_letter(c)}={cell.value!r}|f={fill}")
        if row_vals:
            lines.append(f"  R{r}: " + " || ".join(row_vals))

    # drawing anchors
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if "/drawings/drawing" in n and n.endswith(".xml") and "_rels" not in n:
                xml = z.read(n).decode("utf-8", "replace")
                rows = re.findall(r"<xdr:row>(\d+)</xdr:row>", xml)
                embeds = re.findall(r'r:embed="(rId\d+)"', xml)
                lines.append(f"  drawing {n}: from_rows={rows} embeds={embeds}")
            if "/drawings/_rels/" in n:
                lines.append(f"  rels {n}: {z.read(n).decode('utf-8','replace')}")

OUT.write_text("\n".join(lines), encoding="utf-8")
print(OUT)
print("\n".join(lines[:120]))
