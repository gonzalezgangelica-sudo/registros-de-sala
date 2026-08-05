#!/usr/bin/env python3
"""
Genera plantillas .xlsx (sin macros) a partir de definiciones JSON.
Abiertas en Excel Android / Excel escritorio / Excel Online.
Ideal cuando hay varios registros distintos (p. ej. 7) sin crear una app por cada uno.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
DEFINICIONES = ROOT / "definiciones"
SALIDA = ROOT / "salida"

FILL_LABEL = PatternFill("solid", fgColor="E7E6E6")
FILL_HIGHLIGHT = PatternFill("solid", fgColor="FFF200")
FILL_HEADER = PatternFill("solid", fgColor="F2F2F2")
FONT_BOLD = Font(name="Calibri", size=11, bold=True)
FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="4A4A4A")
FONT_META = Font(name="Calibri", size=10, color="4A4A4A")
FONT_NORMAL = Font(name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="7A7A7A"),
    right=Side(style="thin", color="7A7A7A"),
    top=Side(style="thin", color="7A7A7A"),
    bottom=Side(style="thin", color="7A7A7A"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def safe_filename(text: str) -> str:
    text = re.sub(r"[^\w\-]+", "_", text.strip(), flags=re.UNICODE)
    return text.strip("_") or "registro"


def write_listas(wb: Workbook, listas: dict[str, list[str]]) -> dict[str, str]:
    """Crea hoja LISTAS oculta-ish y devuelve mapa listKey -> rango absoluto."""
    ws = wb.create_sheet("LISTAS")
    named_ranges: dict[str, str] = {}
    col = 1
    for key, values in listas.items():
        letter = get_column_letter(col)
        ws.cell(1, col, key.upper()).font = FONT_BOLD
        ws.cell(1, col).fill = FILL_LABEL
        for i, val in enumerate(values, start=2):
            ws.cell(i, col, val)
        end_row = max(2, 1 + len(values))
        # named range for validation
        ref = f"LISTAS!${letter}$2:${letter}${end_row}"
        named_ranges[key] = ref
        col += 1
    ws.sheet_state = "hidden"
    return named_ranges


def write_meta_header(ws, definicion: dict, hoja: dict, start_row: int = 1) -> int:
    """Cabecera tipo formulario (texto, sin depender de imágenes)."""
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 3, end_column=2)
    cell = ws.cell(start_row, 1, "Stolt Sea Farm")
    cell.font = Font(name="Calibri", size=16, bold=True, color="4A4A4A")
    cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row, end_column=6)
    t = ws.cell(start_row, 3, "FORMULARIO")
    t.font = FONT_TITLE
    t.alignment = CENTER

    ws.merge_cells(start_row=start_row + 1, start_column=3, end_row=start_row + 3, end_column=6)
    sub = ws.cell(start_row + 1, 3, definicion.get("nombre", "").upper())
    sub.font = Font(name="Calibri", size=9, bold=True, color="4A4A4A")
    sub.alignment = CENTER

    meta = [
        ("Código:", definicion.get("codigo", "")),
        ("Revisión:", definicion.get("revision", "")),
        ("Fecha:", definicion.get("fechaDoc", "")),
        ("Hoja:", hoja.get("hojaMeta", "")),
    ]
    for i, (lab, val) in enumerate(meta):
        c1 = ws.cell(start_row + i, 7, lab)
        c2 = ws.cell(start_row + i, 8, val)
        c1.font = FONT_META
        c2.font = FONT_BOLD
        c1.alignment = CENTER
        c2.alignment = CENTER
        c1.border = THIN
        c2.border = THIN
        c1.fill = FILL_HEADER

    for r in range(start_row, start_row + 4):
        for c in range(1, 9):
            ws.cell(r, c).border = THIN

    return start_row + 5


def build_register_sheet(
    wb: Workbook,
    definicion: dict,
    hoja: dict,
    named_lists: dict[str, str],
) -> None:
    ws = wb.create_sheet(hoja["nombre"][:31])
    cols = int(definicion.get("columnas", 8))
    row = write_meta_header(ws, definicion, hoja, 1)

    ws.cell(row, 1, hoja.get("tituloLinea", "")).font = FONT_BOLD
    row += 1

    # header numbers
    header_row = row
    corner = ws.cell(header_row, 1, "")
    corner.fill = FILL_LABEL
    corner.border = THIN
    for i in range(1, cols + 1):
        cell = ws.cell(header_row, i + 1, i)
        cell.font = FONT_BOLD
        cell.alignment = CENTER
        cell.border = THIN

    # fields
    for offset, campo in enumerate(hoja["campos"], start=1):
        r = header_row + offset
        label = ws.cell(r, 1, campo["label"])
        label.font = FONT_BOLD
        label.alignment = CENTER
        label.border = THIN
        label.fill = FILL_HIGHLIGHT if campo.get("highlight") else FILL_LABEL
        ws.row_dimensions[r].height = 28 if "\n" in campo["label"] else 22

        for c in range(2, cols + 2):
            cell = ws.cell(r, c, "")
            cell.border = THIN
            cell.alignment = CENTER
            cell.font = FONT_NORMAL
            if campo.get("type") == "date":
                cell.number_format = "DD/MM/YYYY"

        if campo.get("type") == "select" and campo.get("list") in named_lists:
            ref = named_lists[campo["list"]]
            dv = DataValidation(
                type="list",
                formula1=f"={ref}",
                allow_blank=True,
                showDropDown=False,
            )
            dv.error = "Elige un valor de la lista"
            dv.errorTitle = "Valor no válido"
            rango = f"{get_column_letter(2)}{r}:{get_column_letter(cols + 1)}{r}"
            dv.add(rango)
            ws.add_data_validation(dv)

    # notes
    notes_row = header_row + len(hoja["campos"]) + 2
    ws.cell(notes_row, 1, "Notas / pie de página").font = FONT_BOLD
    for i, note in enumerate(definicion.get("notasPie", [])):
        cell = ws.cell(notes_row + 1 + i, 1, note)
        cell.font = FONT_META
        ws.merge_cells(
            start_row=notes_row + 1 + i,
            start_column=1,
            end_row=notes_row + 1 + i,
            end_column=min(6, cols + 1),
        )
        if i == len(definicion.get("notasPie", [])) - 1:
            cell.fill = FILL_HIGHLIGHT

    # widths
    ws.column_dimensions["A"].width = 18
    for i in range(2, cols + 2):
        ws.column_dimensions[get_column_letter(i)].width = 14

    ws.print_title_rows = "1:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def generate_one(path: Path, out_dir: Path) -> Path:
    data = json.loads(path.read_text(encoding="utf-8"))
    wb = Workbook()
    # remove default
    default = wb.active
    wb.remove(default)

    named = write_listas(wb, data.get("listas", {}))
    for hoja in data.get("hojas", []):
        build_register_sheet(wb, data, hoja, named)

    # instrucciones
    ws = wb.create_sheet("INSTRUCCIONES", 0)
    lines = [
        f"Registro: {data.get('nombre')}",
        f"Código: {data.get('codigo')}  |  Rev: {data.get('revision')}",
        "",
        "Uso en tablet Android:",
        "1. Abre este archivo con Excel (app Microsoft Excel).",
        "2. Rellena las columnas 1–8 (desplegables donde existan).",
        "3. Guarda en OneDrive / SharePoint / carpeta compartida.",
        "4. Opcional: Archivo → Exportar / Imprimir → PDF.",
        "",
        "Sin macros: funciona en Excel Android, Excel Online y escritorio.",
        "Para otro registro distinto: copia un JSON en plantillas/definiciones/ y regenera.",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(i, 1, line).font = FONT_NORMAL if i > 3 else FONT_BOLD
    ws.column_dimensions["A"].width = 100

    out_dir.mkdir(parents=True, exist_ok=True)
    out_name = f"{safe_filename(data.get('id') or path.stem)}.xlsx"
    out_path = out_dir / out_name
    wb.save(out_path)
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera plantillas xlsx de registros de sala")
    parser.add_argument(
        "--all",
        action="store_true",
        help="Genera todas las definiciones (excepto las que empiezan por _)",
    )
    parser.add_argument(
        "--include-examples",
        action="store_true",
        help="Incluye JSON que empiezan por _",
    )
    parser.add_argument("files", nargs="*", help="JSON concretos a generar")
    args = parser.parse_args()

    if args.files:
        sources = [Path(f) for f in args.files]
    else:
        sources = sorted(DEFINICIONES.glob("*.json"))
        if not args.include_examples:
            sources = [p for p in sources if not p.name.startswith("_")]

    if not sources:
        raise SystemExit(f"No hay definiciones en {DEFINICIONES}")

    print(f"Salida: {SALIDA}")
    for src in sources:
        out = generate_one(src, SALIDA)
        print(f"  OK  {src.name} -> {out.name}")


if __name__ == "__main__":
    main()
