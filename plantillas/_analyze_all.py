"""Analiza varios .xlsm de registros de sala: hojas, celdas, VBA."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

FILES = [
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\comprobación detector de metales.xlsm",
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Supervisión del estado de los palés R2.xlsm",
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Registro_LavaUtiles.xlsm",
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Registro de marcas propias .xlsm",
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Registro de marcas propias V2 .xlsm",
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Supervisión de la calidad del transporte R4.xlsm",
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Control de cuchillos produccion R3.xlsm",
    r"k:\Spain\Food Operations\4. Food Plant\Requisitos sala\Nuevos\Control de cutters R4.xlsm",
]

OUT = Path(r"C:\Users\ACZ\.codex\REGISTROS_DE_SALA\plantillas\_analisis")
OUT.mkdir(parents=True, exist_ok=True)


def dump_sheet_preview(ws, max_row=60, max_col=16):
    rows = []
    for r in range(1, min(ws.max_row or 1, max_row) + 1):
        cells = []
        for c in range(1, min(ws.max_column or 1, max_col) + 1):
            v = ws.cell(r, c).value
            if v is not None:
                cells.append({"c": c, "v": str(v)[:200]})
        if cells:
            rows.append({"r": r, "cells": cells})
    return rows


def extract_vba_text(path: Path) -> str:
    try:
        from oletools.olevba import VBA_Parser

        vba = VBA_Parser(str(path))
        parts = []
        if vba.detect_vba_macros():
            for _, _, vba_filename, code in vba.extract_macros():
                if code and code.strip() and code.strip() != "(empty macro)":
                    parts.append(f"===== {vba_filename} =====\n{code}")
        vba.close()
        return "\n\n".join(parts)
    except Exception as e:
        return f"VBA_ERROR: {e}"


def analyze(path_str: str) -> dict:
    path = Path(path_str)
    info = {
        "file": path.name,
        "exists": path.exists(),
        "size": path.stat().st_size if path.exists() else 0,
        "sheets": [],
        "vba_subs": [],
    }
    if not path.exists():
        return info

    wb = load_workbook(path, keep_vba=True, data_only=False)
    for ws in wb.worksheets:
        info["sheets"].append(
            {
                "name": ws.title,
                "state": ws.sheet_state,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "preview": dump_sheet_preview(ws),
            }
        )

    vba = extract_vba_text(path)
    safe = re.sub(r"[^\w\-]+", "_", path.stem)[:60]
    (OUT / f"{safe}_vba.txt").write_text(vba, encoding="utf-8", errors="replace")
    info["vba_subs"] = re.findall(r"(?:Sub|Function)\s+(\w+)", vba)
    info["vba_has_outlook"] = "Outlook" in vba
    info["vba_has_pdf"] = "ExportAsFixedFormat" in vba or "xlTypePDF" in vba
    # save compact json without huge previews duplicated
    (OUT / f"{safe}.json").write_text(
        json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return info


def main():
    summary = []
    for f in FILES:
        print("Analizando:", f)
        info = analyze(f)
        summary.append(
            {
                "file": info["file"],
                "exists": info["exists"],
                "sheets": [s["name"] for s in info.get("sheets", [])],
                "vba_subs": info.get("vba_subs", []),
                "outlook": info.get("vba_has_outlook"),
                "pdf": info.get("vba_has_pdf"),
            }
        )
        print("  sheets:", summary[-1]["sheets"])
        print("  vba:", summary[-1]["vba_subs"])
    (OUT / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("DONE ->", OUT)


if __name__ == "__main__":
    main()
